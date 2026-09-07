"""ISFAHAAN — Intercos : état des dettes intra-groupe à une date, réciprocité, anomalies,
plan de régularisation (virements) — LECTURE SEULE.

Règle cible du groupe : entre deux sociétés SŒURS il ne doit subsister que des dettes
COMMERCIALES (411/401). Toute autre position (455 compte courant, 451 groupe, 467 tiers,
47x virements en transit, prêts) doit être portée par ISFAHAAN (holding) : apports en
compte courant d'associé ou convention de trésorerie, dans un sens ou dans l'autre.

Étapes : 1) balances Pennylane de chaque société à la date d'arrêté ; 2) classification
des comptes interco par LIBELLÉ (convention du cabinet : « <Société> #GISF #GJ », alias
configurables) et par NATURE (cca / tiers / transit / commercial / pret) ; 3) matrice N×N
(ce que A dit que B lui doit) ; 4) RÉCIPROCITÉ : A→B doit valoir −(B→A) ; 5) anomalies ;
6) plan de virements en 3 phases : remontées vers ISFAHAAN (excédents), descentes
ISFAHAAN → sociétés à purger (apport C/C), purges sœur → sœur.
"""
import io
import json
import re
import time
import unicodedata
from collections import defaultdict
from datetime import datetime, timedelta

import httpx
from sqlmodel import Session, select

from app.core.db import engine
from app.core.connectors import pennylane
from app.models import Setting
from . import config
from .treso import _trial_balance

_TZ = timedelta(hours=4)
HOLDING = "ISFAHAAN"

# alias de libellés -> code société (ordre : les plus spécifiques d'abord). Surcharge : Setting ISFAHAAN « interco:aliases ».
DEFAULT_ALIASES = {
    "JBIBFOOD":     ["JB & IB FOOD", "JB&IB FOOD", "JB IB FOOD", "JBIB", "OTC LE PORT", "OTC SAINT-ANDRE", "OTC SAINT ANDRE", "OTC ST ANDRE"],
    "JBFOOD":       ["JB FOOD", "JBFOOD", "OTC SAINT-DENIS", "OTC SAINT DENIS", "OTC ST DENIS"],
    "OTCSTPIERRE":  ["OTC SAINT-PIERRE", "OTC SAINT PIERRE", "OTC ST PIERRE", "OTCSP"],
    "OTCRESERVE":   ["OTC LA RESERVE", "OTCLR"],
    "OTCBRASFUSIL": ["OTC BRAS-FUSIL", "OTC BRAS FUSIL", "BRAS FUSIL", "OTCBF"],
    "GLDSTDENIS":   ["GLD SAINT-DENIS", "GLD SAINT DENIS", "GLD ST DENIS", "GLDSD", "GLDSTD"],
    "GLDCASABONA":  ["GLD CASABONA", "CASABONA"],
    "GONGCHA":      ["GONG CHA", "GONGCHA"],
    "JAG":          ["JAG GLD CHAUDRON", "GLD CHAUDRON", "ENGEN CHAUDRON", "STATION ENGEN CHAUDRON", "SARL JAG", "JAG"],
    "LACORP":       ["LACORP", "LA CORP", "LE COMPTOIR"],
    "ISFAHAAN":     ["ISFAHAAN"],
}
# entités liées connues mais HORS des dossiers Pennylane (affichées à part, jamais dans la matrice)
DEFAULT_EXTERNES = ["CROUSTY GAME", "FPO FAC", "OTC FAC", "GLD LE PORT", "FORMA'CORP", "FORMACORP", "GLD INVEST", "STATION ENGEN"]
TAGS = re.compile(r"#\s*G(ISF|J)\b|\bGISF\b|\bGJ\b")


def _clean(s):
    s = unicodedata.normalize("NFKD", str(s or ""))
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    s = TAGS.sub(" ", s)
    return re.sub(r"\s+", " ", s).strip()


def _nature(number):
    n = str(number)
    if n.startswith(("455", "451")):
        return "cca"
    if n.startswith(("411", "401")):
        return "commercial"
    if n.startswith("47"):
        return "transit"
    if n.startswith(("16", "17", "27")):
        return "pret"
    if n.startswith("46"):
        return "tiers"
    return "autre"


def _settings(code):
    out = {}
    with Session(engine) as s:
        for st in s.exec(select(Setting).where(Setting.company_code == code, Setting.key.in_(["interco:aliases", "interco:externes", "interco:params"]))).all():
            try:
                out[st.key] = json.loads(st.value)
            except Exception:
                pass
    return out


def _balances(pl, day, prefixes=("45", "46", "47", "411", "401", "16", "17", "27"), label_ok=None):
    """Soldes (débit − crédit) des comptes concernés au soir de `day`. trial_balance si le scope
    est ouvert, sinon repli lent : lignes d'écritures par compte."""
    try:
        items = _trial_balance(pl, day)
        return [(str(i["number"]), i.get("label") or "", round(float(i["debits"]) - float(i["credits"]), 2)) for i in items
                if str(i["number"]).startswith(prefixes)], "trial_balance"
    except httpx.HTTPStatusError as e:
        if e.response.status_code != 403:
            raise
    out, cur = [], None
    accs = []
    while True:
        params = {"limit": 100}
        if cur:
            params["cursor"] = cur
        d = pl.get("/ledger_accounts", **params)
        accs += [a for a in (d.get("items") or []) if str(a.get("number", "")).startswith(prefixes)]
        if not d.get("has_more"):
            break
        cur = d.get("next_cursor")
    # repli : seulement les comptes potentiellement interco (45/46/47/51/53, ou libellé taggé/alias)
    accs = [a for a in accs if str(a.get("number", ""))[:2] in ("45", "46", "47", "51", "53")
            or TAGS.search(str(a.get("label") or "").upper()) or (label_ok and label_ok(a.get("label") or ""))]
    for a in accs:
        try:
            lines = pl.account_lines(a["id"], date_to=day)
        except TypeError:
            lines = [l for l in pl.account_lines(a["id"]) if (l.get("date") or "") <= day]
        sol = round(sum(float(l.get("debit") or 0) - float(l.get("credit") or 0) for l in lines if (l.get("date") or "") <= day), 2)
        if abs(sol) >= 0.005:
            out.append((str(a["number"]), a.get("label") or "", sol))
        time.sleep(0.05)
    return out, "lignes (token sans scope balance)"


def _aliases_and_params():
    st = _settings(HOLDING)
    aliases = {k: [_clean(x) for x in v] for k, v in {**DEFAULT_ALIASES, **st.get("interco:aliases", {})}.items()}
    externes = [_clean(x) for x in st.get("interco:externes", DEFAULT_EXTERNES)]
    params = st.get("interco:params", {})
    return aliases, externes, params


def _who(label, me, aliases):
    L = _clean(label)
    for c, al in aliases.items():
        if c != me and any(a and a in L for a in al):
            return c
    return None


def collect(ctx, codes, day, aliases):
    """Balances utiles de chaque société au soir de `day` : (lignes brutes, trésorerie, source)."""
    raw, cash, src = {}, {}, {}
    for k, code in enumerate(codes):
        ctx.progress(k, len(codes) + 1, step=f"balance {code} au {day}…")
        pl = pennylane.for_company(code)
        bal, how = _balances(pl, day, prefixes=("45", "46", "47", "411", "401", "16", "17", "27", "51", "53"),
                             label_ok=lambda lbl, code=code: _who(lbl, code, aliases) is not None)
        raw[code] = bal
        src[code] = how
        cash[code] = round(sum(sol for n, l, sol in bal if n[:2] in ("51", "53")), 2)
    return raw, cash, src


def classify(codes, raw, aliases, externes):
    """Lignes interco classées : (societe, compte, libelle, nature, contrepartie, externe, tagged, solde)."""
    lines = []
    for code in codes:
        for n, lbl, sol in raw[code]:
            if n[:2] in ("51", "53") or abs(sol) < 0.005:
                continue
            nat = _nature(n)
            cp = _who(lbl, code, aliases)
            L = _clean(lbl)
            ext = next((e for e in externes if e in L), None)
            tagged = bool(TAGS.search(str(lbl).upper()))
            collectif = bool(re.fullmatch(r"(455|467|451)0*", n)) or (n.startswith("455") and re.search(r"ASSOCI", L) and not cp)
            perso = nat == "cca" and not cp and re.search(r"C/C|JAFFAR|NASSIR|PRINCIPAL|BLOQU", L)
            if cp is None and not tagged and not ext and not collectif and not perso:
                continue                               # 46x/47x ordinaires (caisse, charges à payer, débiteurs divers…)
            lines.append({"societe": code, "compte": n, "libelle": lbl, "nature": nat, "contrepartie": cp,
                          "externe": ext, "tagged": tagged, "collectif": collectif, "perso": bool(perso), "solde": sol})
    return lines


def analyse(codes, lines, cash, day, floor):
    M = defaultdict(lambda: defaultdict(float))
    Mn = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))
    for l in lines:
        if l["contrepartie"]:
            M[l["societe"]][l["contrepartie"]] += l["solde"]
            Mn[l["societe"]][l["contrepartie"]][l["nature"]] += l["solde"]
    # réciprocité PAR NATURE : hors commercial (455/467/47/prêts = ce que la règle groupe vise) et commercial
    # (411/401 : invérifiable quand une société loge la contrepartie dans un 401/411 collectif, ex. « 401ZEOP »)
    def _nc(d):
        return round(sum(v for k, v in d.items() if k != "commercial"), 2)

    def _co(d):
        return round(d.get("commercial", 0.0), 2)

    recip = []
    for i, a in enumerate(codes):
        for b in codes[i + 1:]:
            ab, ba = round(M[a][b], 2), round(M[b][a], 2)
            if abs(ab) < 0.005 and abs(ba) < 0.005:
                continue
            nc_a, nc_b = _nc(Mn[a][b]), _nc(Mn[b][a])
            co_a, co_b = _co(Mn[a][b]), _co(Mn[b][a])
            ecart_nc = round(nc_a + nc_b, 2)
            co_verif = bool(Mn[a][b].get("commercial") is not None and Mn[b][a].get("commercial") is not None)
            ecart_co = round(co_a + co_b, 2) if co_verif else None
            recip.append({"a": a, "b": b, "a_dit": ab, "b_dit": ba, "ecart": round(ab + ba, 2),
                          "nc_a": nc_a, "nc_b": nc_b, "ecart_nc": ecart_nc, "ok": abs(ecart_nc) < 1.0,
                          "co_a": co_a, "co_b": co_b, "ecart_co": ecart_co, "co_verifiable": co_verif,
                          "a_nat": dict(Mn[a][b]), "b_nat": dict(Mn[b][a])})
    anomalies = []
    for r in recip:
        if not r["ok"]:
            anomalies.append({"type": "non_reciprocite_hors_commercial", "gravite": "haute", "societes": f"{r['a']} ↔ {r['b']}", "montant": r["ecart_nc"],
                              "detail": f"hors 411/401 : {r['a']} voit {r['nc_a']:+,.2f} ; {r['b']} voit {r['nc_b']:+,.2f} (attendu opposés) — écart {r['ecart_nc']:+,.2f}"})
        if r["co_verifiable"] and abs(r["ecart_co"]) >= 1.0:
            anomalies.append({"type": "non_reciprocite_commerciale", "gravite": "moyenne", "societes": f"{r['a']} ↔ {r['b']}", "montant": r["ecart_co"],
                              "detail": f"411/401 : {r['a']} voit {r['co_a']:+,.2f} ; {r['b']} voit {r['co_b']:+,.2f} — écart {r['ecart_co']:+,.2f} (factures non comptabilisées d'un côté / décalage)"})
        elif not r["co_verifiable"] and (abs(r["co_a"]) >= 1.0 or abs(r["co_b"]) >= 1.0):
            anomalies.append({"type": "commercial_non_verifiable", "gravite": "info", "societes": f"{r['a']} ↔ {r['b']}", "montant": r["co_a"] or r["co_b"],
                              "detail": "une des deux sociétés n'a pas de compte 411/401 nominatif pour l'autre (compte collectif) : réciprocité commerciale invérifiable"})
    for a in codes:
        for b in codes:
            if a == b or HOLDING in (a, b):
                continue
            nc = sum(v for k, v in Mn[a][b].items() if k != "commercial")
            if abs(nc) >= 1.0:
                anomalies.append({"type": "dette_non_commerciale_entre_soeurs", "gravite": "haute", "societes": f"{a} → {b}", "montant": round(nc, 2),
                                  "detail": f"{a} porte {nc:+,.2f} sur {b} hors 411/401 ({', '.join(f'{k} {v:+,.0f}' for k, v in Mn[a][b].items() if k != 'commercial')}) — à porter par {HOLDING}"})
    for l in lines:
        if l["nature"] == "transit" and l["contrepartie"] and abs(l["solde"]) >= 1.0:
            anomalies.append({"type": "virement_en_transit", "gravite": "moyenne", "societes": f"{l['societe']} → {l['contrepartie']}", "montant": l["solde"],
                              "detail": f"{l['compte']} « {l['libelle']} » : {l['solde']:+,.2f} non soldé au {day}"})
        if not l["contrepartie"] and l["externe"]:
            anomalies.append({"type": "entite_liee_hors_perimetre", "gravite": "info", "societes": f"{l['societe']} → {l['externe']}", "montant": l["solde"],
                              "detail": f"{l['compte']} « {l['libelle']} » : {l['solde']:+,.2f}"})
        if not l["contrepartie"] and l["tagged"] and not l["externe"]:
            anomalies.append({"type": "interco_non_identifie", "gravite": "moyenne", "societes": l["societe"], "montant": l["solde"],
                              "detail": f"{l['compte']} « {l['libelle']} » : {l['solde']:+,.2f} — libellé taggé groupe mais société inconnue (ajouter un alias)"})
        if not l["contrepartie"] and l.get("collectif") and abs(l["solde"]) >= 1000:
            anomalies.append({"type": "compte_collectif_non_individualise", "gravite": "haute", "societes": l["societe"], "montant": l["solde"],
                              "detail": f"{l['compte']} « {l['libelle']} » : {l['solde']:+,.2f} — contrepartie non identifiable, réciprocité invérifiable (à éclater par société)"})
        if l.get("perso") and abs(l["solde"]) >= 1000:
            anomalies.append({"type": "compte_courant_personne_physique", "gravite": "info", "societes": l["societe"], "montant": l["solde"],
                              "detail": f"{l['compte']} « {l['libelle']} » : {l['solde']:+,.2f} (hors périmètre groupe, pour information)"})

    # ---- plan : dettes nettes non commerciales entre sœurs ----
    pair_debt = {}
    for i, a in enumerate(codes):
        for b in codes[i + 1:]:
            if HOLDING in (a, b):
                continue
            ab = sum(v for k, v in Mn[a][b].items() if k != "commercial")
            ba = sum(v for k, v in Mn[b][a].items() if k != "commercial")
            net = round((ab - ba) / 2, 2) if abs(ab + ba) >= 1.0 else round(ab, 2)
            if abs(net) >= 1.0:
                cred, deb = (a, b) if net > 0 else (b, a)
                pair_debt[(deb, cred)] = abs(net)
    need = defaultdict(float)
    for (deb, cred), amt in pair_debt.items():
        need[deb] += amt
    avail = {c: max(0.0, cash.get(c, 0.0) - floor) for c in codes if c != HOLDING}
    plan, n = [], 0
    descentes = {d: round(max(0.0, need[d] - avail.get(d, 0.0)), 2) for d in need}
    total_desc = sum(descentes.values())
    to_raise = max(0.0, total_desc - max(0.0, cash.get(HOLDING, 0.0) - floor))
    for c in sorted(avail, key=lambda x: -avail[x]):
        if to_raise <= 0:
            break
        surplus = round(avail[c] - need.get(c, 0.0), 2)
        if surplus <= 0:
            continue
        amt = round(min(surplus, to_raise), 2); n += 1
        plan.append({"ordre": n, "phase": "1 · remontée vers la holding", "de": c, "vers": HOLDING, "montant": amt,
                     "libelle": f"CONVENTION TRESORERIE / REMBT C/C ISFAHAAN - {c} - arrete {day}",
                     "condition": f"trésorerie {c} au {day} : {cash.get(c, 0):,.0f} € (plancher {floor:,.0f} €)"})
        to_raise -= amt
    for d, amt in sorted(descentes.items(), key=lambda x: -x[1]):
        if amt <= 0:
            continue
        n += 1
        plan.append({"ordre": n, "phase": "2 · apport C/C holding → société", "de": HOLDING, "vers": d, "montant": amt,
                     "libelle": f"APPORT C/C ASSOCIE ISFAHAAN - {d} - purge intercos {day}",
                     "condition": f"{d} doit {need[d]:,.0f} € à ses sœurs et dispose de {avail.get(d, 0):,.0f} € au-dessus du plancher"})
    for (deb, cred), amt in sorted(pair_debt.items(), key=lambda x: -x[1]):
        n += 1
        plan.append({"ordre": n, "phase": "3 · purge sœur → sœur", "de": deb, "vers": cred, "montant": round(amt, 2),
                     "libelle": f"REGUL INTERCO {deb} -> {cred} - solde au {day}",
                     "condition": "après la phase 2 le cas échéant ; lettrer ensuite les 467/455 réciproques"})
    return {"day": day, "floor": floor, "codes": codes, "cash": cash,
            "matrix": {a: {b: round(M[a][b], 2) for b in codes} for a in codes},
            "matrix_nat": {a: {b: {k: round(v, 2) for k, v in Mn[a][b].items()} for b in codes} for a in codes},
            "lines": lines, "recip": recip, "anomalies": anomalies, "plan": plan, "pair_debt": {f"{d}>{c}": v for (d, c), v in pair_debt.items()},
            "generated": (datetime.utcnow() + _TZ).strftime("%d/%m/%Y %H:%M")}


def run_interco(ctx, day=None, floor=25000.0):
    day = day or ((datetime.utcnow() + _TZ).date().replace(day=1) - timedelta(days=1)).isoformat()
    aliases, externes, params = _aliases_and_params()
    floor = float(params.get("floor", floor))
    codes = [c for c in config.COMPANIES if pennylane.for_company(c)]
    ctx.log(f"Intercos groupe ISFAHAAN au {day} — {len(codes)} sociétés — plancher {floor:,.0f} €")
    raw, cash, src = collect(ctx, codes, day, aliases)
    res = analyse(codes, classify(codes, raw, aliases, externes), cash, day, floor)
    res["sources"] = src
    return res


def excel(result):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    A = "Arial"; white = Font(name=A, bold=True, color="FFFFFF", size=10); navy = PatternFill("solid", fgColor="0A2540")
    base = Font(name=A, size=10); bold = Font(name=A, size=10, bold=True); thin = Border(bottom=Side(style="thin", color="DDDDDD"))
    red = PatternFill("solid", fgColor="F8D2D5"); green = PatternFill("solid", fgColor="D8F0DE"); yellow = PatternFill("solid", fgColor="FFF3BF"); grey = PatternFill("solid", fgColor="EFEFEF")
    codes = result["codes"]; wb = Workbook()

    def head(ws, cols, widths):
        for i, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(row=1, column=i, value=h); c.font = white; c.fill = navy; c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "B2"

    ws = wb.active; ws.title = "Matrice"
    ws["A1"] = f"INTERCOS GROUPE ISFAHAAN — arrêté au {result['day']} — généré le {result['generated']} (Réunion)"; ws["A1"].font = Font(name=A, bold=True, size=13)
    ws["A2"] = "Lecture : la LIGNE est la société dont on lit les livres ; la cellule = ce que cette société dit que la COLONNE lui doit (positif) ou qu'elle doit à la colonne (négatif). Hors 411/401 = dettes non commerciales. Rouge = réciprocité fausse."
    ws["A2"].font = Font(name=A, size=9, color="555555"); ws.merge_cells("A2:L2"); ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 30
    r0 = 4
    ws.cell(row=r0, column=1, value="livres de ↓ / doit à ↓").font = bold
    for j, b in enumerate(codes, 2):
        c = ws.cell(row=r0, column=j, value=b); c.font = white; c.fill = navy; c.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = 15
    ws.column_dimensions["A"].width = 18
    recip_bad = {(r["a"], r["b"]) for r in result["recip"] if not r["ok"]} | {(r["b"], r["a"]) for r in result["recip"] if not r["ok"]}
    for i, a in enumerate(codes, 1):
        ws.cell(row=r0 + i, column=1, value=a).font = bold
        for j, b in enumerate(codes, 2):
            v = result["matrix"][a][b]
            c = ws.cell(row=r0 + i, column=j, value=(v if a != b else None)); c.font = base; c.number_format = "#,##0;[Red]-#,##0;-"
            if a == b: c.fill = grey
            elif (a, b) in recip_bad: c.fill = red
            elif abs(v) >= 1: c.fill = green
    r = r0 + len(codes) + 2
    ws.cell(row=r, column=1, value="Trésorerie (51+53) au " + result["day"]).font = bold
    for j, b in enumerate(codes, 2):
        c = ws.cell(row=r, column=j, value=result["cash"].get(b)); c.number_format = "#,##0"; c.font = base
    ws.cell(row=r + 1, column=1, value="Source").font = bold
    for j, b in enumerate(codes, 2):
        ws.cell(row=r + 1, column=j, value=result["sources"].get(b)).font = Font(name=A, size=8, color="777777")

    w2 = wb.create_sheet("Réciprocité")
    head(w2, ["Société A", "Société B", "HORS COMMERCIAL : A dit", "HORS COMMERCIAL : B dit", "Écart hors commercial", "Réciproque ?",
              "COMMERCIAL : A dit", "COMMERCIAL : B dit", "Écart commercial", "Natures côté A", "Natures côté B"], [14, 14, 18, 18, 16, 11, 16, 16, 16, 34, 34])
    for i, x in enumerate(sorted(result["recip"], key=lambda x: -abs(x["ecart_nc"])), 2):
        vals = [x["a"], x["b"], x["nc_a"], x["nc_b"], x["ecart_nc"], "OUI" if x["ok"] else "NON",
                x["co_a"], x["co_b"], (x["ecart_co"] if x["co_verifiable"] else "invérifiable"),
                ", ".join(f"{k} {v:+,.0f}" for k, v in x["a_nat"].items()), ", ".join(f"{k} {v:+,.0f}" for k, v in x["b_nat"].items())]
        for j, v in enumerate(vals, 1):
            c = w2.cell(row=i, column=j, value=v); c.font = base; c.border = thin
            if j in (3, 4, 5, 7, 8, 9) and isinstance(v, (int, float)): c.number_format = "#,##0.00;[Red]-#,##0.00"
            if j == 6: c.fill = green if x["ok"] else red; c.font = bold

    w3 = wb.create_sheet("Anomalies")
    head(w3, ["Gravité", "Type", "Sociétés", "Montant", "Détail"], [9, 30, 24, 14, 110])
    order = {"haute": 0, "moyenne": 1, "info": 2}
    for i, x in enumerate(sorted(result["anomalies"], key=lambda x: (order[x["gravite"]], -abs(x["montant"] or 0))), 2):
        for j, v in enumerate([x["gravite"], x["type"], x["societes"], x["montant"], x["detail"]], 1):
            c = w3.cell(row=i, column=j, value=v); c.font = base; c.border = thin; c.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 4: c.number_format = "#,##0.00;[Red]-#,##0.00"
            if j == 1: c.fill = red if x["gravite"] == "haute" else (yellow if x["gravite"] == "moyenne" else grey)

    w4 = wb.create_sheet("Plan de virements")
    w4["A1"] = f"FEUILLE D'INSTRUCTION — régularisation des intercos au {result['day']} — plancher de trésorerie {result['floor']:,.0f} € par société"; w4["A1"].font = Font(name=A, bold=True, size=12)
    w4["A2"] = "Exécuter dans l'ORDRE. Phase 1 : les sociétés excédentaires remontent vers ISFAHAAN (convention de trésorerie / remboursement de C/C). Phase 2 : ISFAHAAN apporte en C/C aux sociétés qui doivent purger. Phase 3 : chaque société débitrice règle sa sœur, puis les comptes 467/455 réciproques sont lettrés."
    w4["A2"].font = Font(name=A, size=9, color="555555"); w4.merge_cells("A2:H2"); w4["A2"].alignment = Alignment(wrap_text=True); w4.row_dimensions[2].height = 40
    cols = ["Ordre", "Phase", "De (émetteur)", "Vers (bénéficiaire)", "Montant €", "Libellé du virement", "Condition / contrôle", "Fait (OUI)"]
    for i, (h, w) in enumerate(zip(cols, [7, 30, 16, 18, 14, 52, 60, 10]), 1):
        c = w4.cell(row=4, column=i, value=h); c.font = white; c.fill = navy; w4.column_dimensions[get_column_letter(i)].width = w
    for i, x in enumerate(result["plan"], 5):
        for j, v in enumerate([x["ordre"], x["phase"], x["de"], x["vers"], x["montant"], x["libelle"], x["condition"], ""], 1):
            c = w4.cell(row=i, column=j, value=v); c.font = base; c.border = thin; c.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 5: c.number_format = "#,##0.00"
            if j == 8: c.fill = yellow
    if not result["plan"]:
        w4.cell(row=5, column=1, value="Aucun virement à prévoir : pas de dette non commerciale entre sœurs.").font = base
    w4.freeze_panes = "A5"

    w5 = wb.create_sheet("Détail comptes")
    head(w5, ["Société", "Compte", "Libellé", "Nature", "Contrepartie", "Entité externe", "Solde (débit +)"], [14, 12, 40, 11, 14, 16, 16])
    for i, l in enumerate(sorted(result["lines"], key=lambda l: (l["societe"], -abs(l["solde"]))), 2):
        for j, v in enumerate([l["societe"], l["compte"], l["libelle"], l["nature"], l["contrepartie"] or "", l["externe"] or "", l["solde"]], 1):
            c = w5.cell(row=i, column=j, value=v); c.font = base; c.border = thin
            if j == 7: c.number_format = "#,##0.00;[Red]-#,##0.00"
            if j == 5 and not l["contrepartie"]: c.fill = yellow
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


def run_interco_job(ctx, day=None, floor=25000.0):
    res = run_interco(ctx, day=day, floor=floor)
    stamp = datetime.utcnow() + _TZ
    ctx.add_artifact("json", f"interco_{res['day']}.json", json.dumps(res, ensure_ascii=False).encode("utf-8"), "application/json")
    ctx.add_artifact("xlsx", f"{stamp.strftime('%Y%m%d %H%M')} intercos_groupe_ISFAHAAN_{res['day']} T{ctx.run_id}.xlsx",
                     excel(res), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    hi = sum(1 for a in res["anomalies"] if a["gravite"] == "haute")
    bad = sum(1 for r in res["recip"] if not r["ok"])
    L = [f"INTERCOS GROUPE ISFAHAAN — arrêté {res['day']} — tâche #{ctx.run_id} — {res['generated']}", "",
         f"  sociétés : {len(res['codes'])} · paires en relation : {len(res['recip'])} · réciprocité HORS COMMERCIAL fausse : {bad}",
         f"  anomalies : {len(res['anomalies'])} dont {hi} de gravité haute · virements proposés : {len(res['plan'])}", "",
         "== RÉCIPROCITÉ (écarts) =="] + \
        [f"  {r['a']:13s} ↔ {r['b']:13s} hors commercial : A dit {r['nc_a']:>14,.2f} · B dit {r['nc_b']:>14,.2f} · écart {r['ecart_nc']:>14,.2f}"
         for r in sorted(res["recip"], key=lambda x: -abs(x["ecart_nc"])) if not r["ok"]] + \
        ["", "== PLAN DE VIREMENTS =="] + [f"  {p['ordre']:>2}. [{p['phase']}] {p['de']} → {p['vers']} : {p['montant']:,.2f} € — {p['libelle']}" for p in res["plan"]] + \
        ["", "Excel joint : Matrice · Réciprocité · Anomalies · Plan de virements · Détail comptes."]
    ctx.set_report("\n".join(L))
    return f"{'✅' if not bad and not hi else '⚠️'} Intercos au {res['day']} — {bad} réciprocité(s) fausse(s), {hi} anomalie(s) haute(s), {len(res['plan'])} virement(s) proposé(s)"
