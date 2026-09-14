"""Lecture d'un Board Management Report LP (xlsm) : P&L mensuel par magasin (feuilles 001-ZHY…,
Head Office) et balances mensuelles par entité (feuilles JZ202607 / LBL202607).

P&L : ligne 3 = dates des mois (colonnes F, H, J…) ; lignes repérées par leur libellé (col A/B).
Balance : colonnes = code, libellé, débit/crédit d'ouverture, de période, cumulés, de clôture ;
on ne garde que les lignes de synthèse (pas les sous-lignes analytiques « [001.01]… »)."""
import re
from datetime import datetime, date
from typing import Dict, Tuple

import openpyxl

from . import config

PL_ROWS = [
    ("revenue",         lambda a, b: a.startswith("5101")),
    ("food",            lambda a, b: b.startswith("Operation Food Cost")),
    ("labor",           lambda a, b: b.startswith("Operation Labor Cost")),
    ("tax_ops",         lambda a, b: a.startswith("5402")),
    ("opex_5501",       lambda a, b: a.startswith("5501")),
    ("marketing",       lambda a, b: b.startswith("5501.02")),
    ("utilities",       lambda a, b: b.startswith("5501.06")),
    ("rent",            lambda a, b: b.startswith("5501.09")),
    ("delivery",        lambda a, b: b.startswith("5501.12")),
    ("amort",           lambda a, b: b.startswith("5501.14")),
    ("depr",            lambda a, b: b.startswith("5501.15")),
    ("ga",              lambda a, b: a.startswith("5502")),
    ("fin",             lambda a, b: a.startswith("5503")),
    ("income_tax",      lambda a, b: a.startswith("5701")),
    ("ebitda_report",   lambda a, b: a.startswith("EBITDA (Store)") or a.startswith("EBITDA (Office)")),
    ("ebitda_after_ga", lambda a, b: a.startswith("EBITDA (After G&A)")),
]

BS_CODES = {"1001": "cash_on_hand", "1002": "cash", "1131": "ar", "1133.01": "deposits", "1133.04": "interco_recv",
            "1133.06": "prepaid", "1211": "inventory_food", "1243": "inventory_other", "1501": "fixed_assets_1501", "1901": "lt_prepaid_1901", "2101": "loans", "2121": "ap",
            "2151": "wages_payable", "2171": "tax_payable", "2181.03": "other_03", "2181.12": "interco_pay",
            "2181.13": "advances", "2181.14": "cca_14"}


def _s(v) -> str:
    return str(v).strip() if v is not None else ""


def _num(v) -> float:
    return float(v) if isinstance(v, (int, float)) else 0.0


def parse_pl(ws) -> Dict[str, dict]:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        return {}
    cols = {}
    for j, v in enumerate(rows[2]):
        if isinstance(v, (datetime, date)) and j >= 5:
            cols[j] = f"{v.year:04d}-{v.month:02d}"
    out = {m: {} for m in cols.values()}
    found = set()
    for r in rows:
        a, b = _s(r[0]), _s(r[1])
        for key, test in PL_ROWS:
            if key in found or not test(a, b):
                continue
            found.add(key)
            for j, m in cols.items():
                out[m][key] = _num(r[j]) if j < len(r) else 0.0
    return out


def parse_bs(ws) -> dict:
    out = {}
    for r in ws.iter_rows(values_only=True):
        code, name = _s(r[0]), _s(r[1])
        if code in BS_CODES and not name.startswith("["):
            key = BS_CODES[code]
            if key in out:
                continue
            out[key] = _num(r[8] if len(r) > 8 else 0) - _num(r[9] if len(r) > 9 else 0)   # solde débiteur (+) / créditeur (−)
    out["fixed_assets"] = out.pop("fixed_assets_1501", 0.0) + out.pop("lt_prepaid_1901", 0.0)   # immobilisations brutes + aménagements (net)
    return out


def store_for(sheet: str, month: str) -> str:
    code = config.SHEET_MAP.get(sheet)
    if sheet == "004-QPLFS" and month < config.SHEET_004_SWITCH:
        return "HSF"
    return code


def extract(path_or_stream) -> Tuple[dict, dict, list]:
    """-> (pl: {store: {month: {...}}}, bs: {entity: {month: {...}}}, notes)."""
    wb = openpyxl.load_workbook(path_or_stream, data_only=True, read_only=True)
    pl, bs, notes = {}, {}, []
    for ws in wb.worksheets:
        t = ws.title
        if t in config.SHEET_MAP:
            data = parse_pl(ws)
            for m, v in data.items():
                st = store_for(t, m)
                if st == "HO":
                    if not v.get("ga"):
                        continue
                elif not v.get("revenue") and not v.get("labor"):
                    continue          # mois vide (magasin pas encore ouvert / fermé)
                pl.setdefault(st, {})[m] = v
        mm = re.fullmatch(r"(JZ|LBL)(\d{4})(\d{2})", t)
        if mm:
            bs.setdefault(mm.group(1), {})[f"{mm.group(2)}-{mm.group(3)}"] = parse_bs(ws)
    if not pl and not bs:
        notes.append("Aucune feuille reconnue (attendu : 001-ZHY, 002-BFC, 004-QPLFS, Head Office, JZ2026MM, LBL2026MM…).")
    return pl, bs, notes
