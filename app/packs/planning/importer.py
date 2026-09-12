"""Import des exports Skello (classeur mensuel « Détails » + contrats) : salariés, compétences déduites,
habitudes (jours non travaillés, CFA, dimanches), plages réalisées."""
import json
import re
import unicodedata
import collections
import datetime as dt
from typing import List, Dict

import openpyxl
from sqlmodel import Session, select

from app.core.db import engine
from app.models import PlEmployee, PlShift
from . import config, service

# libellé Skello -> clé de poste
_ALIASES = {
    "VENTE MATIN": "VENTE_MATIN", "VENTE APRES MIDI": "VENTE_APREM", "VENTE APRES-MIDI": "VENTE_APREM", "VENTE JOURNEE": "VENTE_JOURNEE",
    "VENTE RESPONSABLE": "VENTE_RESP", "VENTE RESPONSABLE ADJOINT E": "VENTE_RESP", "VENDEUR APPRENTI E": "VENDEUR_APP", "VENTE STAGIAIRE": "VENDEUR_APP",
    "BOULANGERIE 3H": "BOULANGERIE_3H", "BOULANGERIE 5H": "BOULANGERIE_5H", "BOULANGERIE 10H": "BOULANGERIE_10H",
    "BOULANGERIE APPRENTI E": "BOULANGERIE_APP", "BOULANGERIE STAGIAIRE": "BOULANGERIE_STAG", "BOULANGERIE": "BOULANGERIE_5H",
    "PATISSERIE 5H": "PATISSERIE_5H", "PATISSERIE 05H00": "PATISSERIE_5H", "PATISSERIE 5H10": "PATISSERIE_5H", "PATISSERIE 7H": "PATISSERIE_7H",
    "PATISSERIE 9H": "PATISSERIE_7H", "PATISSERIE RESPONSABLE": "PATISSERIE_RESP", "PATISSERIE STAGIAIRE": "PATISSERIE_STAG", "PATISSERIE APPRENTI E": "PATISSERIE_STAG",
    "TRAITEUR": "TRAITEUR", "TRAITEUR RESPONSABLE": "TRAITEUR_RESP", "TRAITEUR APPRENTI E": "TRAITEUR_APP", "TRATEUR STAGIAIRE": "TRAITEUR_APP",
    "ADMINISTRATIF": "ADMINISTRATIF", "RESP EXPLOITATION": "ADMINISTRATIF", "AGENT POLYVALENT": "AGENT_POLYVALENT", "INVENTAIRE": "INVENTAIRE",
}
_ABS = {"repos hebdomadaire": "Repos hebdomadaire", "congé payé": "Congé payé", "arrêt maladie": "Arrêt maladie", "école - cfa": "École - CFA",
        "jour férié": "Jour férié", "récupération": "Récupération", "absence injustifiée": "Absence injustifiée", "absence autorisée": "Absence autorisée",
        "repos compensateur": "Repos compensateur"}


def _norm(label: str) -> str:
    t = unicodedata.normalize("NFKD", str(label or "")).encode("ascii", "ignore").decode().upper()
    t = re.sub(r"\s*-\s*(COPAIN|KOOKA|KOOK).*$", "", t)          # « VENTE Matin - Copain La… » -> « VENTE MATIN »
    t = re.sub(r"[^A-Z0-9]+", " ", t).strip()
    return t


def post_key_of(label: str) -> str:
    n = _norm(label)
    if n in _ALIASES:
        return _ALIASES[n]
    if n.startswith("FORMATION") or "FORMATION" in n:
        return "FORMATION"
    for k, v in _ALIASES.items():
        if n.startswith(k):
            return v
    return service._norm_key(n)


def absence_type_of(label: str) -> str:
    l = str(label or "").strip().lower()
    for k, v in _ABS.items():
        if l.startswith(k):
            return v
    if "maladie" in l:
        return "Arrêt maladie"
    if "cfa" in l or "école" in l:
        return "École - CFA"
    if "férié" in l:
        return "Jour férié"
    return str(label or "Absence")[:40]


def _f(v, default=0.0) -> float:
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "", "-", "None") else default
    except (TypeError, ValueError):
        return default


def read_details(paths: List[str]) -> List[dict]:
    rows = []
    for p in paths:
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = wb["Détails"]
        it = ws.iter_rows(values_only=True)
        hdr = [str(h) for h in next(it)]
        col = lambda pre: next((h for h in hdr if h.startswith(pre)), None)
        C = {"post": col("Nom du poste"), "site": col("Shift / abs"), "ht": col("Heures Travaill"), "abs_in": col("Absences Incluses"),
             "abs_out": col("Absences non incl"), "nuit": col("Heures de Nuit")}
        for r in it:
            if not any(r):
                continue
            d = dict(zip(hdr, r))
            dd = d.get("Date")
            if not isinstance(dd, (dt.date, dt.datetime)):
                continue
            rows.append({
                "first": str(d.get("Prénom") or "").strip(), "last": str(d.get("Nom") or "").strip(),
                "contract": str(d.get("Type de contrat") or "CDI"), "hours": _f(d.get("Temps contractuel"), 35.0) if str(d.get("Heures / jours") or "heures").startswith("heure") else 35.0,
                "site": str(d.get("Etablissement principal") or ""), "shift_site": str(d.get(C["site"]) or ""),
                "date": dd.date() if isinstance(dd, dt.datetime) else dd, "kind": str(d.get("Travail / Absence") or ""),
                "label": str(d.get(C["post"]) or ""), "start": d.get("Début"), "end": d.get("Fin"),
                "pause": _f(d.get("Pause (h)")), "ht": _f(d.get(C["ht"])),
                "abs_h": _f(d.get(C["abs_in"])) + _f(d.get(C["abs_out"])),
                "note": str(d.get("Note") or "")[:200],
            })
    return rows


def read_contracts(paths: List[str]) -> Dict[str, dict]:
    """Premier onglet (« CONTRATS STANDARDS ») : dates de début/fin et intitulé de poste, dernier classeur gagnant."""
    out = {}
    for p in sorted(paths):
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        for r in ws.iter_rows(values_only=True):
            if not r or r[0] in (None, "Nom", "CONTRATS STANDARDS", "EXTRAS", "") or r[1] in (None, "Prénom"):
                continue
            try:
                key = f"{str(r[1]).strip()} {str(r[0]).strip()}".upper()
                def pd(x):
                    if isinstance(x, (dt.date, dt.datetime)):
                        return x.date() if isinstance(x, dt.datetime) else x
                    if x and re.match(r"\d{2}/\d{2}/\d{2}$", str(x)):
                        return dt.datetime.strptime(str(x), "%d/%m/%y").date()
                    return None
                out[key] = {"hours": _f(r[2], 35.0), "contract": str(r[3] or "CDI"), "title": str(r[4] or ""), "start": pd(r[5]), "end": pd(r[6])}
            except Exception:
                continue
    return out


def site_code(name: str) -> str:
    n = _norm(name)
    if "STLEU" in n or "ST LEU" in n or "SAINT LEU" in n:
        return "SL"
    if "POSSESSION" in n:
        return "LP"
    if "STEMARIE" in n or "STE MARIE" in n or "SAINTE MARIE" in n:
        return "SM"
    return "SL"


def deduce_coverage(rows: List[dict], site: str) -> dict:
    """Médiane du nombre de personnes par poste et jour de semaine, par saison, sur les plages réalisées du site."""
    import statistics
    W = [r for r in rows if r["kind"] == "travail" and site_code(r["shift_site"]) == site and 0.5 < r["ht"] <= 12]
    seasons = {"default": None, "dec_jan": (12, 1), "jul_aug": (7, 8), "mar_mai": (3, 4, 5)}
    out = {}
    for sk, months in seasons.items():
        rs = [r for r in W if months is None or r["date"].month in months]
        wk = sorted({r["date"].isocalendar()[:2] for r in rs})
        if len(wk) < 4:
            continue
        cov = {}
        for pk in {post_key_of(r["label"]) for r in rs}:
            if pk == "FORMATION":
                continue
            med = []
            for wd in range(7):
                med.append(int(statistics.median([sum(1 for r in rs if post_key_of(r["label"]) == pk and r["date"].isocalendar()[:2] == w and r["date"].weekday() == wd) for w in wk])))
            if any(med):
                cov[pk] = med
        out[sk] = cov
    return out


def import_skello(paths: List[str], company_code: str = config.COMPANY_CODE, site: str = None,
                  shifts_from: dt.date = None, shifts_to: dt.date = None, active_since: dt.date = None) -> dict:
    """Crée/actualise les salariés de TOUS les établissements présents dans les exports (pool société : établissement
    principal = celui où la personne a le plus de plages), déduit compétences et habitudes, importe les plages réalisées
    d'une période sur tous les sites, et déduit la couverture des sites qui n'en ont pas encore."""
    rows = read_details(paths)
    contracts = read_contracts(paths)
    dates = [r["date"] for r in rows]
    last = max(dates)
    active_since = active_since or (last - dt.timedelta(days=60))
    byemp = collections.defaultdict(list)
    seen = set()
    for r in rows:
        k = (r["first"].upper(), r["last"].upper(), r["date"], r["kind"], r["label"], str(r["start"]), str(r["end"]), site_code(r["shift_site"]))
        if k in seen:
            continue                     # même ligne présente dans deux exports (salarié multi-sites)
        seen.add(k)
        byemp[f"{r['first']} {r['last']}".upper()].append(r)
    for st in {site_code(r["shift_site"]) for r in rows if r["kind"] == "travail"}:
        service.ensure_posts(company_code, st)
    created = updated = 0
    with Session(engine) as s:
        existing = {e.external_key: e for e in s.exec(select(PlEmployee).where(PlEmployee.company_code == company_code)).all() if e.external_key}
        for key, rs in byemp.items():
            work = [r for r in rs if r["kind"] == "travail" and r["ht"] > 0.5 and r["ht"] <= 12]
            if not work:
                continue
            first, lastn = rs[0]["first"], rs[0]["last"]
            ct = contracts.get(key, {})
            recent = [r for r in work if r["date"] >= active_since]
            active = bool(recent) and (not ct.get("end") or ct["end"] >= last)
            # compétences : part des plages par poste
            cnt = collections.Counter(post_key_of(r["label"]) for r in work if post_key_of(r["label"]) != "FORMATION")
            tot = sum(cnt.values()) or 1
            posts = {}
            for k, c in cnt.items():
                share = c / tot
                if share >= 0.35:
                    posts[k] = 1                       # préféré (poste principal)
                elif share >= 0.12 or c >= 15:
                    posts[k] = 2                       # tient le poste
                elif c >= 5:
                    posts[k] = 3                       # peut dépanner
            # habitudes
            wdays = {r["date"] for r in work}
            weeks = len({d.isocalendar()[:2] for d in wdays}) or 1
            pres = {wd: round(100 * sum(1 for d in wdays if d.weekday() == wd) / weeks) for wd in range(7)}
            days_off = [wd for wd in range(7) if weeks >= 8 and pres[wd] < 12]
            cfa = collections.Counter(r["date"].weekday() for r in rs if "cfa" in r["label"].lower())
            cfa_days = [wd for wd, c in cfa.items() if c >= 5 and c >= 0.6 * sum(cfa.values())]
            sunday = "oui" if pres[6] >= 8 else "non"
            hours = ct.get("hours") or rs[-1]["hours"]
            max_days = 5 if hours >= 30 else (4 if hours >= 25 else (3 if hours >= 20 else 2))
            main_site = collections.Counter(site_code(r["shift_site"]) for r in work).most_common(1)[0][0]
            other = sorted({site_code(r["shift_site"]) for r in work if site_code(r["shift_site"]) != main_site})
            fields = dict(site=main_site, first_name=first, last_name=lastn, contract_type=ct.get("contract") or rs[-1]["contract"], weekly_hours=hours,
                          start_date=ct.get("start"), end_date=ct.get("end"), active=active, posts=json.dumps(posts), days_off=json.dumps(days_off),
                          cfa_days=json.dumps(cfa_days), pattern=json.dumps(pres), sunday=sunday, max_days=max_days, mobility=json.dumps(other),
                          source="skello", external_key=key)
            e = existing.get(key)
            if e:
                for k, v in fields.items():
                    setattr(e, k, v)
                updated += 1
            else:
                e = PlEmployee(company_code=company_code, **fields)
                created += 1
            s.add(e)
        s.commit()
        emap = {e.external_key: e.id for e in s.exec(select(PlEmployee).where(PlEmployee.company_code == company_code)).all() if e.external_key}
        # plages réalisées (tous sites)
        n_sh = 0
        if shifts_from and shifts_to:
            for x in s.exec(select(PlShift).where(PlShift.company_code == company_code, PlShift.source == "import",
                                                   PlShift.date >= shifts_from, PlShift.date <= shifts_to)).all():
                s.delete(x)
            for r in [x for v in byemp.values() for x in v]:
                if not (shifts_from <= r["date"] <= shifts_to):
                    continue
                eid = emap.get(f"{r['first']} {r['last']}".upper())
                if not eid:
                    continue
                if r["kind"] == "travail":
                    if not isinstance(r["start"], dt.datetime) or not isinstance(r["end"], dt.datetime) or r["ht"] <= 0.02:
                        continue
                    st, en = r["start"].strftime("%H:%M"), r["end"].strftime("%H:%M")
                    key = post_key_of(r["label"])
                    kind = "task" if key == "FORMATION" or "LIVRAISON" in _norm(r["label"]) else "work"
                    s.add(PlShift(company_code=company_code, site=site_code(r["shift_site"]), employee_id=eid, date=r["date"], kind=kind, post_key=key if kind == "work" else r["label"][:40],
                                  start=st, end=en, pause=r["pause"], hours=round(r["ht"], 2), note=r["note"] or None, status="published", source="import"))
                else:
                    typ = absence_type_of(r["label"])
                    if typ == "Repos hebdomadaire":
                        continue                     # le repos est l'absence de plage
                    s.add(PlShift(company_code=company_code, site=site_code(r["site"]), employee_id=eid, date=r["date"], kind="absence", post_key=typ,
                                  hours=round(r["abs_h"], 2), status="published", source="import"))
                n_sh += 1
            s.commit()
    # couverture déduite pour les sites qui n'en ont pas encore (le site principal garde la grille du questionnaire)
    cfg = service.get_config(company_code)
    cbs = dict(cfg.get("coverage_by_site") or {})
    added = []
    main_sites = {e.site for e in service.employees(company_code, None, active_only=False)}
    for st in {site_code(r["shift_site"]) for r in rows if r["kind"] == "travail"}:
        if st != cfg.get("site") and st not in cbs and st in main_sites:   # pas de grille déduite pour un site sans effectif propre
            cov = deduce_coverage([x for v in byemp.values() for x in v], st)
            if cov.get("default"):
                cbs[st] = cov
                added.append(st)
    if added:
        service.save_config({"coverage_by_site": cbs}, company_code)
    return {"employees_created": created, "employees_updated": updated, "shifts": n_sh, "period": (min(dates), last), "coverage_added": added}
